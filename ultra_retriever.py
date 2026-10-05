import asyncio
import os
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np


class HungarianTextNormalizer:

    STOPWORDS = {
        "a", "az", "és", "hogy", "van", "volt", "lesz", "egy", "is", "nem",
        "vagy", "mint", "de", "ha", "mi", "ki", "hol", "mikor", "mely",
    }
    SUFFIXES = re.compile(
        r"(nak|nek|ban|ben|ról|ről|tól|től|hoz|hez|höz|val|vel|ba|be|ra|re|"
        r"ok|ek|ök|ak|at|et|ot|öt)$",
        re.IGNORECASE,
    )

    @classmethod
    def normalize(cls, text: str) -> List[str]:
        tokens = re.findall(r"\b[\wáéíóöőúüű]+\b", text.lower())
        return [
            cls.SUFFIXES.sub("", token) if len(token) > 5 else token
            for token in tokens
            if token not in cls.STOPWORDS and len(token) > 1
        ]


@dataclass(frozen=True)
class RetrievedDocument:
    chunk_id: str
    source: str
    text: str


class UltraHybridRetriever:
    SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}
    CHUNK_SIZE = 1200
    CHUNK_OVERLAP = 150

    def __init__(
        self,
        vector_db_client=None,
        embedder=None,
        database_path: Optional[str] = None,
        username: Optional[str] = None,
        documents_dir: Optional[str] = None,
        min_dense_similarity: float = 0.28,
    ):
        self.db = vector_db_client
        self.embedder = embedder
        self.username = username or os.getenv("ZOLI_USERNAME")
        self.database_path = Path(
            database_path or os.getenv("ZOLI_DATABASE_PATH", "zoli_gpt_enterprise.db")
        )
        configured_dir = documents_dir or os.getenv("ZOLI_DOCUMENTS_DIR")
        self.documents_dir = Path(configured_dir).expanduser() if configured_dir else None
        self.min_dense_similarity = min_dense_similarity
        self.normalizer = HungarianTextNormalizer()
        self.documents = self._load_documents()
        self._tokenized = [self.normalizer.normalize(doc.text) for doc in self.documents]
        self._document_vectors = self._encode_documents()

    @property
    def document_count(self) -> int:
        return len(self.documents)

    def _load_documents(self) -> List[RetrievedDocument]:
        documents = self._load_database_documents()
        documents.extend(self._load_directory_documents())
        unique_documents = []
        seen = set()
        for document in documents:
            key = (document.source, document.text)
            if document.text.strip() and key not in seen:
                seen.add(key)
                unique_documents.append(document)
        return unique_documents[:10000]

    def _load_database_documents(self) -> List[RetrievedDocument]:
        if not self.username or not self.database_path.is_file():
            return []

        try:
            connection = sqlite3.connect(
                f"{self.database_path.resolve().as_uri()}?mode=ro", uri=True
            )
            columns = {
                row[1]
                for row in connection.execute("PRAGMA table_info(document_vectors)")
            }
            required = {"id", "username", "doc_name", "chunk_text"}
            if not required.issubset(columns):
                connection.close()
                return []
            rows = connection.execute(
                "SELECT id, doc_name, chunk_text FROM document_vectors "
                "WHERE username = ? ORDER BY id LIMIT 10000",
                (self.username,),
            ).fetchall()
            connection.close()
        except (OSError, sqlite3.Error):
            return []

        return [
            RetrievedDocument(f"db:{row_id}", str(doc_name), str(text))
            for row_id, doc_name, text in rows
            if text and str(text).strip()
        ]

    def _load_directory_documents(self) -> List[RetrievedDocument]:
        root = self.documents_dir
        if not root or not root.is_dir():
            return []

        documents = []
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in self.SUPPORTED_SUFFIXES:
                continue
            try:
                if path.stat().st_size > 10 * 1024 * 1024:
                    continue
                text = self._read_file(path)
                source = path.relative_to(root).as_posix()
                for index, chunk in enumerate(self._chunk_text(text), start=1):
                    documents.append(
                        RetrievedDocument(f"file:{source}:{index}", source, chunk)
                    )
            except (OSError, ValueError, ImportError):
                continue
            if len(documents) >= 10000:
                break
        return documents

    @staticmethod
    def _read_file(path: Path) -> str:
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader

            return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        if path.suffix.lower() == ".docx":
            from docx import Document

            return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        return path.read_text(encoding="utf-8", errors="ignore")

    @classmethod
    def _chunk_text(cls, text: str) -> List[str]:
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        chunks = []
        start = 0
        while start < len(text):
            end = min(start + cls.CHUNK_SIZE, len(text))
            if end < len(text):
                boundary = max(
                    text.rfind("\n", start + cls.CHUNK_SIZE // 2, end),
                    text.rfind(". ", start + cls.CHUNK_SIZE // 2, end),
                    text.rfind(" ", start + cls.CHUNK_SIZE // 2, end),
                )
                if boundary > start:
                    end = boundary + 1
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(text):
                break
            start = max(start + 1, end - cls.CHUNK_OVERLAP)
        return chunks

    def _encode(self, texts: List[str]) -> np.ndarray:
        try:
            vectors = self.embedder.encode(
                texts, normalize_embeddings=True, show_progress_bar=False
            )
        except TypeError:
            vectors = self.embedder.encode(texts, normalize_embeddings=True)
        return np.atleast_2d(np.asarray(vectors, dtype=np.float32))

    def _encode_documents(self):
        if not self.embedder or not self.documents:
            return None
        try:
            vectors = self._encode([document.text for document in self.documents])
            if vectors.ndim != 2 or vectors.shape[0] != len(self.documents):
                return None
            norms = np.linalg.norm(vectors, axis=1, keepdims=True)
            return vectors / np.maximum(norms, 1e-12)
        except Exception:
            return None

    def _bm25_scores(self, query_tokens: List[str]) -> List[float]:
        if not query_tokens or not self._tokenized:
            return [0.0] * len(self.documents)

        document_count = len(self._tokenized)
        lengths = [len(tokens) for tokens in self._tokenized]
        average_length = sum(lengths) / max(document_count, 1)
        document_frequency: Dict[str, int] = {}
        for tokens in self._tokenized:
            for token in set(tokens):
                document_frequency[token] = document_frequency.get(token, 0) + 1

        scores = []
        for tokens, length in zip(self._tokenized, lengths):
            frequencies: Dict[str, int] = {}
            for token in tokens:
                frequencies[token] = frequencies.get(token, 0) + 1
            score = 0.0
            for token in set(query_tokens):
                term_frequency = frequencies.get(token, 0)
                if not term_frequency:
                    continue
                inverse_frequency = np.log(
                    1 + (document_count - document_frequency[token] + 0.5)
                    / (document_frequency[token] + 0.5)
                )
                denominator = term_frequency + 1.5 * (
                    1 - 0.75 + 0.75 * length / max(average_length, 1)
                )
                score += inverse_frequency * term_frequency * 2.5 / denominator
            scores.append(float(score))
        return scores

    def search(self, query: str, top_k: int = 5) -> List[RetrievedDocument]:
        if not self.documents or top_k <= 0:
            return []

        sparse_scores = self._bm25_scores(self.normalizer.normalize(query))
        sparse_hits = sorted(
            (index for index, score in enumerate(sparse_scores) if score > 0),
            key=lambda index: sparse_scores[index],
            reverse=True,
        )[:20]

        dense_hits = []
        if self._document_vectors is not None:
            try:
                query_vector = self._encode([query])[0]
                query_vector /= max(float(np.linalg.norm(query_vector)), 1e-12)
                similarities = self._document_vectors @ query_vector
                dense_hits = sorted(
                    (
                        index
                        for index, score in enumerate(similarities)
                        if score >= self.min_dense_similarity
                    ),
                    key=lambda index: float(similarities[index]),
                    reverse=True,
                )[:20]
            except Exception:
                dense_hits = []

        fused_scores: Dict[int, float] = {}
        for hits in (sparse_hits, dense_hits):
            for rank, index in enumerate(hits):
                fused_scores[index] = fused_scores.get(index, 0.0) + 1.0 / (60 + rank + 1)

        ranked_indices = sorted(
            fused_scores, key=lambda index: fused_scores[index], reverse=True
        )[:top_k]
        return [self.documents[index] for index in ranked_indices]

    async def retrieve_and_compress(
        self, query: str, top_k: int = 5, max_chars: int = 6000
    ) -> str:
        documents = await asyncio.to_thread(self.search, query, top_k)
        sections = []
        remaining = max_chars
        for index, document in enumerate(documents, start=1):
            section = f"[L{index}] Forrás: {document.source}\n{document.text}"
            if len(section) > remaining:
                if remaining <= 0:
                    break
                section = section[:remaining]
            sections.append(section)
            remaining -= len(section) + 2
        return "\n\n".join(sections)